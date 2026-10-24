# -*- coding: utf-8 -*-
"""Arma el libro en formato Word (.docx) a partir de los módulos de contenido."""
import os, re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import dsl
import c1_general, c2_173a, c3_173b, c4_174_180, c5_practica, c6_mas, c7_cierre  # noqa: F401 (cargan BLOCKS)
import json

NAVY = RGBColor(0x14, 0x21, 0x3D)
WINE = RGBColor(0xB2, 0x3A, 0x48)
GREY = RGBColor(0x55, 0x5B, 0x66)
BODY_FONT = "Cambria"
HEAD_FONT = "Calibri"
SLIDES = os.path.join(HERE, "slides_png")
OUT = os.path.join(HERE, "Estafas_y_defraudaciones_Codigo_Penal_argentino.docx")

doc = Document()

# ------------------------------------------------------------------ estilos
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
sec.top_margin, sec.bottom_margin = Cm(2.3), Cm(2.0)
sec.left_margin, sec.right_margin = Cm(2.4), Cm(2.2)
sec.header_distance, sec.footer_distance = Cm(1.1), Cm(1.0)

def set_font(style, name, size, bold=None, color=None, italic=None):
    style.font.name = name
    style.font.size = Pt(size)
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)
    if bold is not None: style.font.bold = bold
    if italic is not None: style.font.italic = italic
    if color is not None: style.font.color.rgb = color

st = doc.styles
set_font(st["Normal"], BODY_FONT, 10.5)
pf = st["Normal"].paragraph_format
pf.space_after = Pt(5); pf.space_before = Pt(0); pf.line_spacing = 1.12
pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

for name, size, color, before, after in (("Heading 1", 19, NAVY, 0, 10), ("Heading 2", 14, WINE, 12, 5), ("Heading 3", 11.5, NAVY, 8, 3)):
    s = st[name]
    set_font(s, HEAD_FONT, size, bold=True, color=color, italic=False)
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    s.paragraph_format.line_spacing = 1.0
st["Heading 1"].paragraph_format.page_break_before = True

for lst in ("List Bullet", "List Number"):
    set_font(st[lst], BODY_FONT, 10.5)
    st[lst].paragraph_format.space_after = Pt(3)
    st[lst].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

# ------------------------------------------------------------------ utilidades
TOKEN = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*)")

def add_rich(p, text, size=None, color=None, italic=False, bold=False, font=None):
    for part in TOKEN.split(text):
        if not part:
            continue
        b, i = bold, italic
        if part.startswith("**") and part.endswith("**"):
            part, b = part[2:-2], True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, i = part[1:-1], (not italic)
        r = p.add_run(part)
        r.bold = b; r.italic = i
        if size: r.font.size = Pt(size)
        if color is not None: r.font.color.rgb = color
        if font:
            r.font.name = font
            r._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    return p

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)

def cell_borders(cell, left=None, others="nil", lw=24):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        if side == "left" and left:
            el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(lw)); el.set(qn("w:color"), left)
        else:
            el.set(qn("w:val"), others)
        b.append(el)
    tcPr.append(b)

def cell_margins(cell, top=90, bottom=90, left=160, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{k}"); el.set(qn("w:w"), str(v)); el.set(qn("w:type"), "dxa"); m.append(el)
    tcPr.append(m)

def table_borders(tbl, color="C9CDD6", sz=4):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(sz)); el.set(qn("w:color"), color)
        b.append(el)
    tblPr.append(b)

def no_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit"); trPr.append(el)

def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader"); trPr.append(el)

def set_col_widths(tbl, widths_cm):
    tbl.autofit = False
    tblPr = tbl._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    grid = tbl._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths_cm):
        gc.set(qn("w:w"), str(int(w * 567)))
    for idx, w in enumerate(widths_cm):
        tbl.columns[idx].width = Cm(w)
    for row in tbl.rows:
        for idx, w in enumerate(widths_cm):
            row.cells[idx].width = Cm(w)

CONTENT_W = 16.4  # ancho útil en cm

def keep_next(p):
    p.paragraph_format.keep_with_next = True

def box(label, color_hex, fill_hex, paras, label_color=None):
    """Recuadro de una celda con etiqueta y párrafos (lista de (texto, opciones))."""
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]
    c.width = Cm(CONTENT_W)
    shade(c, fill_hex); cell_borders(c, left=color_hex); cell_margins(c)
    no_split(t.rows[0])
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(label.upper()); r.bold = True; r.font.size = Pt(8.5)
    r.font.color.rgb = label_color or RGBColor.from_string(color_hex); r.font.name = HEAD_FONT
    for txt, opt in paras:
        for line in txt.split("\n"):
            q = c.add_paragraph()
            q.paragraph_format.space_after = Pt(3)
            q.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_rich(q, line, size=opt.get("size", 10), italic=opt.get("italic", False), bold=opt.get("bold", False), color=opt.get("color"))
    spacer(4)

def spacer(pt=4):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(pt)
    r = p.add_run(""); r.font.size = Pt(2)

def add_field(paragraph, instr):
    r = paragraph.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = "Actualice el índice: clic derecho → Actualizar campo."
    f3 = OxmlElement("w:fldChar"); f3.set(qn("w:fldCharType"), "end")
    r._r.append(f1); r._r.append(it); r._r.append(f2); r._r.append(t); r._r.append(f3)

# ------------------------------------------------------------------ portada
def cover():
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(70)
    for i in range(1):
        pass
    t = doc.add_table(rows=1, cols=1); c = t.rows[0].cells[0]
    shade(c, "14213D"); cell_borders(c, left="B23A48", lw=60); cell_margins(c, 500, 500, 500, 400)
    c.width = Cm(CONTENT_W)
    q = c.paragraphs[0]; q.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = q.add_run("CÓDIGO PENAL ARGENTINO · ARTS. 172 A 180"); r.font.size = Pt(10); r.bold = True
    r.font.color.rgb = RGBColor(0xF2, 0xC1, 0x4E); r.font.name = HEAD_FONT
    q = c.add_paragraph(); q.paragraph_format.space_before = Pt(10); q.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = q.add_run("Estafas y otras defraudaciones"); r.font.size = Pt(34); r.bold = True
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); r.font.name = HEAD_FONT
    q = c.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = q.add_run("Apunte de estudio y conferencia: doctrina, jurisprudencia, ejemplos y práctica en Córdoba, con filminas incorporadas")
    r.font.size = Pt(14); r.font.color.rgb = RGBColor(0xCB, 0xD5, 0xE1); r.font.name = HEAD_FONT
    q = c.add_paragraph(); q.paragraph_format.space_before = Pt(40); q.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = q.add_run("Estafa (art. 172) · Casos especiales (art. 173) · Agravantes (art. 174) · Atenuantes y usura (arts. 175 y 175 bis) · Quebrados e insolvencia (arts. 176 a 180) · Cheques (art. 302) · Excusa absolutoria (art. 185)")
    r.font.size = Pt(10.5); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); r.font.name = HEAD_FONT
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(30)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img = os.path.join(SLIDES, "s05.png")
    p.add_run().add_picture(img, width=Cm(13.5))
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(30)
    add_rich(p, "**Texto legal vigente a octubre de 2026**", size=11, color=NAVY)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_rich(p, "Córdoba, República Argentina", size=10, color=GREY)

cover()

# ------------------------------------------------------------------ índice
h = doc.add_paragraph("Índice", style="Heading 1")
PAGES_FILE = os.path.join(HERE, "build", "toc_pages.json")
toc_pages = json.load(open(PAGES_FILE, encoding="utf-8")) if os.path.exists(PAGES_FILE) else {}
toc_items = [(b[0], b[1]) for b in dsl.BLOCKS if b[0] in ("h1", "h2")]

def toc_par(level, text, page, first=False, last=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.space_after = Pt(1 if level == 2 else 2); pf.space_before = Pt(0 if level == 2 else 5)
    pf.left_indent = Cm(0 if level == 1 else 0.6)
    pf.tab_stops.add_tab_stop(Cm(CONTENT_W), alignment=2, leader=1)  # derecha, puntos
    if first:
        r = p.add_run()
        for typ, txt in (("begin", None), ("instr", 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None)):
            if typ == "instr":
                el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = txt
            else:
                el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), typ)
            r._r.append(el)
    r = p.add_run(text); r.font.size = Pt(10 if level == 1 else 9.3); r.bold = (level == 1)
    r.font.name = HEAD_FONT if level == 1 else BODY_FONT
    if level == 1: r.font.color.rgb = NAVY
    r = p.add_run("\t" + str(page)); r.font.size = Pt(10 if level == 1 else 9.3); r.bold = (level == 1)
    if last:
        r = p.add_run(); el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), "end"); r._r.append(el)

for i, (lv, txt) in enumerate(toc_items):
    toc_par(1 if lv == "h1" else 2, txt, toc_pages.get(str(i), "00"), first=(i == 0), last=(i == len(toc_items) - 1))

# ------------------------------------------------------------------ bloques
fig_count = 0
for b in dsl.BLOCKS:
    k = b[0]
    if k == "parte":
        p = doc.add_paragraph(); p.paragraph_format.page_break_before = True
        p.paragraph_format.space_before = Pt(200); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(b[1].upper()); r.font.size = Pt(14); r.bold = True; r.font.color.rgb = WINE; r.font.name = HEAD_FONT
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(b[2]); r.font.size = Pt(26); r.bold = True; r.font.color.rgb = NAVY; r.font.name = HEAD_FONT
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("━━━━━━━━"); r.font.color.rgb = WINE
    elif k == "h1":
        doc.add_paragraph(b[1], style="Heading 1")
    elif k == "h2":
        doc.add_paragraph(b[1], style="Heading 2")
    elif k == "h3":
        doc.add_paragraph(b[1], style="Heading 3")
    elif k == "p":
        add_rich(doc.add_paragraph(), b[1])
    elif k == "bul":
        for it in b[1]:
            add_rich(doc.add_paragraph(style="List Bullet"), it)
    elif k == "num":
        for i, it in enumerate(b[1], 1):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.7); p.paragraph_format.first_line_indent = Cm(-0.5)
            p.paragraph_format.space_after = Pt(3)
            add_rich(p, f"**{i}.** " + it)
    elif k == "ley":
        box("Texto legal · " + b[1], "14213D", "EEF0F4", [(b[2], {"italic": True, "size": 10.2})])
    elif k == "res":
        box("En síntesis", "14213D", "E4EAF5", [(b[1], {"size": 10})])
    elif k == "ej":
        box("Ejemplo · " + b[1], "C9A227", "FBF5E6", [(b[2], {"size": 10})], label_color=RGBColor(0x8A, 0x6A, 0x0B))
    elif k == "jur":
        paras = [("**" + b[1] + "**", {"size": 9.8}), (b[2], {"size": 10})]
        if b[3]:
            paras.append(("Fuente: " + b[3], {"size": 8.8, "italic": True, "color": GREY}))
        box("Jurisprudencia", "B23A48", "F7ECEE", paras)
    elif k == "tip":
        box("Para la práctica", "2E7D5B", "E9F4EE", [(b[1], {"size": 10})])
    elif k == "ojo":
        box("Atención", "C77700", "FFF4DC", [(b[1], {"size": 9.8})], label_color=RGBColor(0x9A, 0x5B, 0x00))
    elif k == "ficha":
        p = doc.add_paragraph(); keep_next(p); p.paragraph_format.space_after = Pt(2)
        r = p.add_run("FICHA TÉCNICA"); r.bold = True; r.font.size = Pt(8.5); r.font.color.rgb = WINE; r.font.name = HEAD_FONT
        t = doc.add_table(rows=0, cols=2); table_borders(t, "D5D9E0")
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for key, val in b[1]:
            row = t.add_row(); no_split(row)
            c0, c1 = row.cells
            shade(c0, "14213D"); cell_margins(c0, 50, 50, 110, 80); cell_margins(c1, 50, 50, 120, 100)
            q = c0.paragraphs[0]; q.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = q.add_run(key); r.bold = True; r.font.size = Pt(9.3); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); r.font.name = HEAD_FONT
            q = c1.paragraphs[0]; q.alignment = WD_ALIGN_PARAGRAPH.LEFT
            q.paragraph_format.space_after = Pt(0)
            add_rich(q, val, size=9.5)
        set_col_widths(t, [3.6, CONTENT_W - 3.6])
        spacer(6)
    elif k == "tab":
        headers, rows, widths = b[1], b[2], b[3]
        t = doc.add_table(rows=1, cols=len(headers)); table_borders(t, "D5D9E0")
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr = t.rows[0]; repeat_header(hdr); no_split(hdr)
        for i, hd in enumerate(headers):
            c = hdr.cells[i]; shade(c, "14213D"); cell_margins(c, 60, 60, 100, 80)
            q = c.paragraphs[0]; q.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = q.add_run(hd); r.bold = True; r.font.size = Pt(9.2); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); r.font.name = HEAD_FONT
        for ri, rw in enumerate(rows):
            row = t.add_row(); no_split(row)
            for i, val in enumerate(rw):
                c = row.cells[i]; cell_margins(c, 50, 50, 100, 80)
                if ri % 2: shade(c, "F6F4EF")
                q = c.paragraphs[0]; q.alignment = WD_ALIGN_PARAGRAPH.LEFT
                q.paragraph_format.space_after = Pt(0); q.paragraph_format.line_spacing = 1.0
                add_rich(q, val, size=9.2, bold=(i == 0), color=(WINE if i == 0 else None))
        widths = widths or [1] * len(headers)
        tot = sum(widths)
        set_col_widths(t, [CONTENT_W * w / tot for w in widths])
        spacer(6)
    elif k == "fig":
        n, cap = b[1], b[2]
        img = os.path.join(SLIDES, f"s{n:02d}.png")
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(2)
        keep_next(p)
        p.add_run().add_picture(img, width=Cm(15.2))
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(8)
        add_rich(p, f"**Filmina {n}.** {cap}", size=8.8, color=GREY, italic=True)
    elif k == "pb":
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    else:
        raise ValueError(k)

# ------------------------------------------------------------------ encabezado y pie
for s in doc.sections:
    s.different_first_page_header_footer = True
    hp = s.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = hp.add_run("Estafas y otras defraudaciones · Código Penal argentino"); r.font.size = Pt(8); r.font.color.rgb = GREY; r.italic = True
    fp = s.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run(); r.font.size = Pt(9)
    for typ, txt in (("begin", None), ("instr", "PAGE"), ("separate", None), ("t", "1"), ("end", None)):
        if typ == "instr":
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = txt
        elif typ == "t":
            el = OxmlElement("w:t"); el.text = txt
        else:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), typ)
        r._r.append(el)

# Idioma español (corrector de Word)
rpr = st["Normal"].element.get_or_add_rPr()
lang = OxmlElement("w:lang"); lang.set(qn("w:val"), "es-AR"); rpr.append(lang)

# Propiedades del documento
cp = doc.core_properties
cp.title = "Estafas y otras defraudaciones en el Código Penal argentino"
cp.subject = "Apunte de estudio y conferencia (arts. 172 a 180 CP)"
cp.language = "es-AR"
cp.keywords = "estafa; defraudación; art. 172; art. 173; Código Penal; Córdoba"

doc.save(OUT)
print("OK", OUT, "bloques:", len(dsl.BLOCKS))
